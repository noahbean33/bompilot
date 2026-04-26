"""
Tests for Item 10: part locking, re-match skip, monitoring skip, manufacturer BOM export.
"""

import csv
import io
import hashlib
import secrets
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.models.part_result import PartResult as PartResultRow
from app.models.project import BomLine, Project
from app.providers.base import ComponentProvider
from app.providers.registry import ProviderRegistry, get_registry
from app.providers.schema import (
    DistributorStock,
    ParametricQuery,
    PartResult as ProviderResult,
    PriceBreak,
    ProviderCapabilities,
)
from app.schemas.preferences import MergedPreferences
from app.models.password_reset_token import PasswordResetToken
from app.services.auth import get_user_by_email
from app.worker.tasks import run_monitor

# ---------------------------------------------------------------------------
# Helpers shared with other test files
# ---------------------------------------------------------------------------

REGISTER_URL = "/auth/register"
LOGIN_URL = "/auth/login"
PROJECTS_URL = "/projects/"

_SAMPLE_CSV = (
    "Reference,Value,Footprint,Description,Quantity,MPN\n"
    "C1,100nF,C_0402,Decoupling cap,10,GRM155R61A104KA01D\n"
    "R1,10k,R_0402,Pull-up,5,RC0402FR-0710KL\n"
)


def _register_and_login(client, db_session, email, password="pass123"):
    # Register with email and name (no password)
    client.post(REGISTER_URL, json={"email": email, "name": "Test User"})

    # Create a password reset token via db_session
    raw_token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
    expires_at = datetime.now(UTC) + timedelta(hours=24)

    user = get_user_by_email(db_session, email=email)
    db_session.add(PasswordResetToken(
        user_id=user.id,
        token_hash=token_hash,
        expires_at=expires_at,
    ))
    db_session.commit()

    # Reset password using the raw token
    client.post("/auth/reset-password", json={"token": raw_token, "new_password": password})

    # Login
    resp = client.post(LOGIN_URL, json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def _create_project(client, token, name="Test"):
    resp = client.post(PROJECTS_URL, json={"name": name}, headers=_auth(token))
    assert resp.status_code == 201
    return resp.json()


def _import_bom(client, project_id, token, csv_text=_SAMPLE_CSV):
    resp = client.post(
        f"/projects/{project_id}/bom/import",
        files={"file": ("bom.csv", io.BytesIO(csv_text.encode()), "text/csv")},
        headers=_auth(token),
    )
    assert resp.status_code == 200


def _get_line(client, project_id, token, ref):
    lines = client.get(f"/projects/{project_id}/bom", headers=_auth(token)).json()
    return next(ln for ln in lines if ln["reference"] == ref)


# ---------------------------------------------------------------------------
# Mock provider
# ---------------------------------------------------------------------------

_FAKE_RESULT = ProviderResult(
    mpn="GRM155R61A104KA01D",
    manufacturer="Murata",
    description="100nF 0402",
    package="0402",
    pricing=[PriceBreak(quantity=1, unit_price=0.10, currency="USD")],
    stock_total=50000,
    distributors=[DistributorStock(distributor="digikey", stock=50000, url=None)],
    lifecycle_status="active",
    tech_specs=None,
    datasheet_url=None,
    similar_parts=None,
    source_provider="mock",
    retrieved_at=datetime(2026, 4, 7, tzinfo=UTC),
    match_type="exact_mpn",
)


class MockProvider(ComponentProvider):
    def __init__(self, results=None):
        self._results = results if results is not None else [_FAKE_RESULT]

    async def search_by_mpn(self, mpn, quantity, preferences):
        return self._results

    async def search_by_distributor_pn(self, distributor, pn, quantity):
        raise NotImplementedError

    async def search_parametric(self, params, preferences):
        raise NotImplementedError

    def capabilities(self):
        return ProviderCapabilities(
            has_lifecycle_status=True,
            has_tech_specs=False,
            has_datasheet_urls=False,
            has_similar_parts=False,
            has_parametric_search=False,
            distributor_coverage=["digikey"],
        )


def _mock_registry(results=None):
    reg = ProviderRegistry()
    reg.register("mock", MockProvider(results))
    reg.set_active("mock")
    return reg


# ---------------------------------------------------------------------------
# Lock / unlock endpoint tests
# ---------------------------------------------------------------------------


def test_lock_sets_locked_true(client: TestClient, db_session: Session):
    token = _register_and_login(client, db_session, "lock1@example.com")
    project = _create_project(client, token)
    pid = project["id"]
    _import_bom(client, pid, token)
    line = _get_line(client, pid, token, "C1")

    resp = client.post(
        f"/projects/{pid}/bom/{line['id']}/lock", headers=_auth(token)
    )
    assert resp.status_code == 200
    assert resp.json()["locked"] is True


def test_unlock_sets_locked_false(client: TestClient, db_session: Session):
    token = _register_and_login(client, db_session, "lock2@example.com")
    project = _create_project(client, token)
    pid = project["id"]
    _import_bom(client, pid, token)
    line = _get_line(client, pid, token, "C1")

    client.post(f"/projects/{pid}/bom/{line['id']}/lock", headers=_auth(token))
    resp = client.post(
        f"/projects/{pid}/bom/{line['id']}/unlock", headers=_auth(token)
    )
    assert resp.status_code == 200
    assert resp.json()["locked"] is False


def test_lock_wrong_user_returns_404(client: TestClient, db_session: Session):
    token_a = _register_and_login(client, db_session, "lock3a@example.com")
    token_b = _register_and_login(client, db_session, "lock3b@example.com")
    project = _create_project(client, token_a)
    pid = project["id"]
    _import_bom(client, pid, token_a)
    line = _get_line(client, pid, token_a, "C1")

    resp = client.post(
        f"/projects/{pid}/bom/{line['id']}/lock", headers=_auth(token_b)
    )
    assert resp.status_code == 404


def test_lock_requires_auth(client: TestClient):
    resp = client.post("/projects/1/bom/1/lock")
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# Re-matching skips locked lines
# ---------------------------------------------------------------------------


def test_rematch_skips_locked_lines(client: TestClient, db_session: Session):
    """Locked lines are not re-matched when POST /match is called."""
    reg = _mock_registry()
    app.dependency_overrides[get_registry] = lambda: reg
    try:
        token = _register_and_login(client, db_session, "lock4@example.com")
        project = _create_project(client, token)
        pid = project["id"]
        _import_bom(client, pid, token)

        # First match — both lines get matched
        client.post(f"/projects/{pid}/match", headers=_auth(token))

        # Lock C1 and record its current selected_result_id
        c1 = _get_line(client, pid, token, "C1")
        c1_id = c1["id"]
        c1_result_id_before = c1["selected_result"]["id"]
        client.post(f"/projects/{pid}/bom/{c1_id}/lock", headers=_auth(token))

        # Re-match with a different provider result
        new_result = _FAKE_RESULT.model_copy(update={"mpn": "DIFFERENT-MPN"})
        app.dependency_overrides[get_registry] = lambda: _mock_registry([new_result])
        client.post(f"/projects/{pid}/match", headers=_auth(token))

        # C1 should still have the old result (locked → skipped)
        c1_after = _get_line(client, pid, token, "C1")
        assert c1_after["selected_result"]["id"] == c1_result_id_before

        # R1 (unlocked) should have been updated
        r1_after = _get_line(client, pid, token, "R1")
        assert r1_after["selected_result"]["mpn"] == "DIFFERENT-MPN"
    finally:
        app.dependency_overrides.pop(get_registry, None)


# ---------------------------------------------------------------------------
# Monitoring skips locked lines
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_monitor_skips_locked_lines(db_session: Session):
    """run_monitor does not re-query provider for locked BOM lines."""
    project = Project(user_id=1, name="p")
    db_session.add(project)
    db_session.flush()

    # Locked BOM line
    locked_line = BomLine(
        project_id=project.id,
        reference="C1",
        mpn_raw="GRM155R61A104KA01D",
        quantity=10,
        raw_fields={},
        locked=True,
    )
    db_session.add(locked_line)

    # Unlocked BOM line
    unlocked_line = BomLine(
        project_id=project.id,
        reference="R1",
        mpn_raw="RC0402",
        quantity=1,
        raw_fields={},
        locked=False,
    )
    db_session.add(unlocked_line)
    db_session.flush()

    locked_result = PartResultRow(
        bom_line_id=locked_line.id,
        rank=1,
        mpn="GRM155R61A104KA01D",
        manufacturer="Murata",
        description="Cap",
        package="0402",
        distributor="digikey",
        unit_price=0.10,
        stock=50000,
        lifecycle_status=None,
        tech_specs=None,
        datasheet_url=None,
        source_provider="mock",
        match_type="exact_mpn",
        retrieved_at=datetime(2026, 4, 7, tzinfo=UTC),
    )
    unlocked_result = PartResultRow(
        bom_line_id=unlocked_line.id,
        rank=1,
        mpn="RC0402",
        manufacturer="Yageo",
        description="Resistor",
        package="0402",
        distributor="digikey",
        unit_price=0.01,
        stock=100000,
        lifecycle_status=None,
        tech_specs=None,
        datasheet_url=None,
        source_provider="mock",
        match_type="exact_mpn",
        retrieved_at=datetime(2026, 4, 7, tzinfo=UTC),
    )
    db_session.add_all([locked_result, unlocked_result])
    db_session.commit()

    # Provider will now return zero stock to trigger a flag
    call_log: list[str] = []

    class TrackingProvider(MockProvider):
        async def search_by_mpn(self, mpn, quantity, preferences):
            call_log.append(mpn)
            zero_stock = _FAKE_RESULT.model_copy(update={"mpn": mpn, "stock_total": 0})
            return [zero_stock]

    reg = ProviderRegistry()
    reg.register("mock", TrackingProvider())
    reg.set_active("mock")

    await run_monitor(db_session, reg)

    # Only the unlocked line's MPN should have been queried
    assert "GRM155R61A104KA01D" not in call_log
    assert "RC0402" in call_log


# ---------------------------------------------------------------------------
# Manufacturer BOM export
# ---------------------------------------------------------------------------


def test_export_manufacturer_csv_columns(client: TestClient, db_session: Session):
    """CSV export contains all required columns in correct order."""
    reg = _mock_registry()
    app.dependency_overrides[get_registry] = lambda: reg
    try:
        token = _register_and_login(client, db_session, "exp1@example.com")
        project = _create_project(client, token, "Test Project")
        pid = project["id"]
        _import_bom(client, pid, token)
        client.post(f"/projects/{pid}/match", headers=_auth(token))

        resp = client.get(
            f"/projects/{pid}/export/manufacturer?format=csv",
            headers=_auth(token),
        )
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        assert "_manufacturer_bom.csv" in resp.headers["content-disposition"]

        # Strip BOM and parse
        text = resp.content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        assert len(rows) == 2  # C1 and R1

        expected_cols = [
            "Reference", "Quantity", "Value", "Footprint", "MPN",
            "Manufacturer", "Description", "Distributor", "Unit Price",
            "Currency", "Stock", "Notes",
        ]
        assert reader.fieldnames == expected_cols
    finally:
        app.dependency_overrides.pop(get_registry, None)


def test_export_manufacturer_csv_unresolved_lines_included(client: TestClient, db_session: Session):
    """Unresolved lines (no match) are included with blank MPN/Manufacturer fields."""
    token = _register_and_login(client, db_session, "exp2@example.com")
    project = _create_project(client, token, "Exp Project")
    pid = project["id"]
    _import_bom(client, pid, token)
    # No matching run — all lines unresolved

    resp = client.get(
        f"/projects/{pid}/export/manufacturer?format=csv",
        headers=_auth(token),
    )
    assert resp.status_code == 200
    text = resp.content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    assert len(rows) == 2  # both lines included
    for row in rows:
        assert row["MPN"] == ""
        assert row["Manufacturer"] == ""
        assert row["Reference"] != ""  # Reference is populated


def test_export_manufacturer_csv_locked_and_unlocked_both_included(client: TestClient, db_session: Session):
    """Both locked and unlocked lines appear in the CSV export."""
    reg = _mock_registry()
    app.dependency_overrides[get_registry] = lambda: reg
    try:
        token = _register_and_login(client, db_session, "exp3@example.com")
        project = _create_project(client, token, "Exp3 Project")
        pid = project["id"]
        _import_bom(client, pid, token)
        client.post(f"/projects/{pid}/match", headers=_auth(token))

        # Lock one line
        c1 = _get_line(client, pid, token, "C1")
        client.post(f"/projects/{pid}/bom/{c1['id']}/lock", headers=_auth(token))

        resp = client.get(
            f"/projects/{pid}/export/manufacturer?format=csv",
            headers=_auth(token),
        )
        assert resp.status_code == 200
        text = resp.content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        assert len(rows) == 2  # both C1 (locked) and R1 (unlocked)
        refs = {r["Reference"] for r in rows}
        assert "C1" in refs
        assert "R1" in refs
    finally:
        app.dependency_overrides.pop(get_registry, None)


def test_export_manufacturer_xlsx_columns(client: TestClient, db_session: Session):
    """Excel export has correct columns and sheet name."""
    import openpyxl

    reg = _mock_registry()
    app.dependency_overrides[get_registry] = lambda: reg
    try:
        token = _register_and_login(client, db_session, "exp4@example.com")
        project = _create_project(client, token, "Exp4 Project")
        pid = project["id"]
        _import_bom(client, pid, token)
        client.post(f"/projects/{pid}/match", headers=_auth(token))

        resp = client.get(
            f"/projects/{pid}/export/manufacturer?format=xlsx",
            headers=_auth(token),
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        assert "_manufacturer_bom.xlsx" in resp.headers["content-disposition"]

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "BOM" in wb.sheetnames
        ws = wb["BOM"]

        header_row = [cell.value for cell in ws[1]]
        expected = [
            "Reference", "Quantity", "Value", "Footprint", "MPN",
            "Manufacturer", "Description", "Distributor", "Unit Price",
            "Currency", "Stock", "Notes",
        ]
        assert header_row == expected
        # Header should be bold
        assert ws.cell(1, 1).font.bold is True
        # Data rows: 2 lines + header = 3 rows total
        assert ws.max_row == 3
    finally:
        app.dependency_overrides.pop(get_registry, None)


def test_export_requires_auth(client: TestClient):
    resp = client.get("/projects/1/export/manufacturer")
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# Bulk lock / unlock
# ---------------------------------------------------------------------------


def test_lock_all_sets_locked_true_on_all_rows(client: TestClient, db_session: Session):
    token = _register_and_login(client, db_session, "bulk_lock1@example.com")
    project = _create_project(client, token)
    pid = project["id"]
    _import_bom(client, pid, token)

    resp = client.post(f"/projects/{pid}/bom/lock-all", headers=_auth(token))
    assert resp.status_code == 200
    assert resp.json()["locked_count"] == 2

    lines = client.get(f"/projects/{pid}/bom", headers=_auth(token)).json()
    assert all(ln["locked"] is True for ln in lines)


def test_unlock_all_sets_locked_false_on_all_rows(client: TestClient, db_session: Session):
    token = _register_and_login(client, db_session, "bulk_lock2@example.com")
    project = _create_project(client, token)
    pid = project["id"]
    _import_bom(client, pid, token)

    # Lock everything first
    client.post(f"/projects/{pid}/bom/lock-all", headers=_auth(token))
    lines_locked = client.get(f"/projects/{pid}/bom", headers=_auth(token)).json()
    assert all(ln["locked"] is True for ln in lines_locked)

    resp = client.post(f"/projects/{pid}/bom/unlock-all", headers=_auth(token))
    assert resp.status_code == 200
    assert resp.json()["unlocked_count"] == 2

    lines = client.get(f"/projects/{pid}/bom", headers=_auth(token)).json()
    assert all(ln["locked"] is False for ln in lines)


def test_lock_all_returns_correct_count(client: TestClient, db_session: Session):
    """locked_count reflects the actual number of rows updated."""
    token = _register_and_login(client, db_session, "bulk_lock3@example.com")
    project = _create_project(client, token)
    pid = project["id"]

    # Import 2-line BOM
    _import_bom(client, pid, token)

    resp = client.post(f"/projects/{pid}/bom/lock-all", headers=_auth(token))
    assert resp.json()["locked_count"] == 2


def test_lock_all_404_wrong_user(client: TestClient, db_session: Session):
    token_a = _register_and_login(client, db_session, "bulk_lock4a@example.com")
    token_b = _register_and_login(client, db_session, "bulk_lock4b@example.com")
    project = _create_project(client, token_a)
    pid = project["id"]

    resp = client.post(f"/projects/{pid}/bom/lock-all", headers=_auth(token_b))
    assert resp.status_code == 404


def test_unlock_all_404_wrong_user(client: TestClient, db_session: Session):
    token_a = _register_and_login(client, db_session, "bulk_lock5a@example.com")
    token_b = _register_and_login(client, db_session, "bulk_lock5b@example.com")
    project = _create_project(client, token_a)
    pid = project["id"]

    resp = client.post(f"/projects/{pid}/bom/unlock-all", headers=_auth(token_b))
    assert resp.status_code == 404


def test_lock_all_requires_auth(client: TestClient):
    resp = client.post("/projects/1/bom/lock-all")
    assert resp.status_code == 401


def test_unlock_all_requires_auth(client: TestClient):
    resp = client.post("/projects/1/bom/unlock-all")
    assert resp.status_code == 401
